#!/usr/bin/env python3
"""Extract content from xlsx and docx files."""
import sys

# Extract xlsx content
print("=" * 80)
print("附件4-技术评分表.xlsx 内容:")
print("=" * 80)
try:
    from openpyxl import load_workbook
    wb = load_workbook('/workspace/docs/附件4-技术评分表.xlsx', data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n--- Sheet: {sheet_name} (rows={ws.max_row}, cols={ws.max_column}) ---")
        for row in ws.iter_rows(values_only=True):
            # Skip empty rows
            if any(cell is not None and str(cell).strip() != '' for cell in row):
                print(" | ".join(str(c) if c is not None else '' for c in row))
except Exception as e:
    print(f"Error reading xlsx: {e}")

# Extract docx content
print("\n\n" + "=" * 80)
print("附件8-测试方案.docx 内容:")
print("=" * 80)
try:
    from docx import Document
    doc = Document('/workspace/docs/附件8-测试方案.docx')
    
    print("\n--- Paragraphs ---")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            print(f"[P{i}] [{para.style.name}] {para.text}")
    
    print("\n--- Tables ---")
    for ti, table in enumerate(doc.tables):
        print(f"\n[Table {ti}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            print(" | ".join(cells))
except Exception as e:
    print(f"Error reading docx: {e}")
